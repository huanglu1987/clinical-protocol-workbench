"""Build the deterministic synthetic DOCX pair used by the S3 Word check."""

from __future__ import annotations

import io
import json
import sys
import zipfile
from copy import copy
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "worker"))

from clinical_qc.docx_package import (  # noqa: E402
    COMMENTS,
    COMMENT_TYPE,
    CT,
    DOC_RELS,
    MAIN,
    R,
    REL,
    W,
    DocxPackage,
    digest,
    serialize,
    w,
)
from clinical_qc.redline import (  # noqa: E402
    AUTHOR,
    Change,
    freeze_plan,
    generate,
    make_anchor,
    visible_text,
)


SOURCE = ROOT / "docs/verification/fixtures/s3/synthetic-native-review-source-v0.1.docx"
OUTPUT = ROOT / "docs/verification/artifacts/s3/synthetic-native-review-redline-v0.1.docx"
EVIDENCE = ROOT / "docs/verification/evidence/2026-09-05-s3-macos"
PLAN = EVIDENCE / "frozen-change-plan.json"
MANIFEST = EVIDENCE / "build-manifest.json"
FIXED_ZIP_TIME = (2026, 9, 5, 2, 0, 0)
REVIEW_TIME = "2026-09-05T02:00:00Z"
OLD_REVIEW_TIME = "2026-08-01T00:00:00Z"
OLD_AUTHOR = "原审阅者"


def _set_font(run, size: float = 11, bold: bool = False) -> None:
    run.font.name = "Songti SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), "Songti SC")


def _base_docx() -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "Songti SC"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run("S3 Word 原生审阅链路合成验证"), size=16, bold=True)
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(notice.add_run("全部内容均为虚构，不得用于临床决策或申报。"), size=10)

    heading = document.add_paragraph()
    _set_font(heading.add_run("1 合成测试正文"), size=13, bold=True)
    _set_font(document.add_paragraph().add_run("完成访视,记录结果。"))
    _set_font(document.add_paragraph().add_run("给药剂量需医学负责人确认。"))
    _set_font(document.add_paragraph().add_run("历史审阅段落占位"))
    _set_font(document.add_paragraph().add_run("已有批注位置"))

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "检查项"
    table.cell(0, 1).text = "预期结果"
    table.cell(1, 0).text = "原生审阅"
    table.cell(1, 1).text = "旧历史保留；新增修订和批注可见"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_font(run, size=10, bold=row is table.rows[0])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run("S3 合成材料 | 仅供内部技术验证"), size=9)

    properties = document.core_properties
    properties.title = "S3 Word 原生审阅链路合成验证"
    properties.subject = "经典批注与简单修订子集"
    properties.author = "Clinical Protocol Workbench Synthetic Fixture"
    properties.created = datetime(2026, 9, 5, 2, 0, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 9, 5, 2, 0, tzinfo=timezone.utc)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _paragraph_with_text(document: etree._Element, value: str) -> etree._Element:
    for paragraph in document.iter(w("p")):
        if "".join(paragraph.xpath(".//w:t/text()", namespaces={"w": W})) == value:
            return paragraph
    raise RuntimeError(f"synthetic paragraph not found: {value}")


def _clear_content(paragraph: etree._Element) -> None:
    for child in list(paragraph):
        if child.tag != w("pPr"):
            paragraph.remove(child)


def _review_run(text: str, *, deleted: bool = False) -> etree._Element:
    run = etree.Element(w("r"))
    text_node = etree.SubElement(run, w("delText") if deleted else w("t"))
    text_node.text = text
    return run


def _inject_existing_review(base: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(base)) as archive:
        infos = [copy(info) for info in archive.infolist()]
        parts = {info.filename: archive.read(info.filename) for info in infos}

    document = etree.fromstring(parts[MAIN])
    history = _paragraph_with_text(document, "历史审阅段落占位")
    _clear_content(history)
    deleted = etree.SubElement(
        history,
        w("del"),
        {w("id"): "7", w("author"): OLD_AUTHOR, w("date"): OLD_REVIEW_TIME},
    )
    deleted.append(_review_run("旧文字", deleted=True))
    inserted = etree.SubElement(
        history,
        w("ins"),
        {w("id"): "8", w("author"): OLD_AUTHOR, w("date"): OLD_REVIEW_TIME},
    )
    inserted.append(_review_run("待接受文字"))

    commented = _paragraph_with_text(document, "已有批注位置")
    original_runs = [child for child in list(commented) if child.tag == w("r")]
    if len(original_runs) != 1:
        raise RuntimeError("unexpected synthetic comment paragraph")
    position = commented.index(original_runs[0])
    commented.insert(position, etree.Element(w("commentRangeStart"), {w("id"): "12"}))
    commented.insert(position + 2, etree.Element(w("commentRangeEnd"), {w("id"): "12"}))
    reference_run = etree.Element(w("r"))
    etree.SubElement(reference_run, w("commentReference"), {w("id"): "12"})
    commented.insert(position + 3, reference_run)
    parts[MAIN] = serialize(document)

    comments = etree.Element(w("comments"), nsmap={"w": W})
    comment = etree.SubElement(
        comments,
        w("comment"),
        {w("id"): "12", w("author"): OLD_AUTHOR, w("date"): OLD_REVIEW_TIME},
    )
    comment_p = etree.SubElement(comment, w("p"))
    comment_r = etree.SubElement(comment_p, w("r"))
    etree.SubElement(comment_r, w("t")).text = "不要删除此批注。"
    parts[COMMENTS] = serialize(comments)

    relationships = etree.fromstring(parts[DOC_RELS])
    used_ids = {item.get("Id") for item in relationships}
    relation_id = next(f"rId{index}" for index in range(1, 1000) if f"rId{index}" not in used_ids)
    etree.SubElement(
        relationships,
        f"{{{REL}}}Relationship",
        Id=relation_id,
        Type=R + "/comments",
        Target="comments.xml",
    )
    parts[DOC_RELS] = etree.tostring(
        relationships.getroottree(), encoding="UTF-8", xml_declaration=True, standalone=True
    )

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(
        content_types,
        f"{{{CT}}}Override",
        PartName="/" + COMMENTS,
        ContentType=COMMENT_TYPE,
    )
    parts["[Content_Types].xml"] = etree.tostring(
        content_types.getroottree(), encoding="UTF-8", xml_declaration=True, standalone=True
    )

    settings = etree.fromstring(parts["word/settings.xml"])
    if settings.find(w("trackRevisions")) is None:
        settings.insert(0, etree.Element(w("trackRevisions")))
    parts["word/settings.xml"] = etree.tostring(
        settings.getroottree(), encoding="UTF-8", xml_declaration=True, standalone=True
    )

    ordered_names = [info.filename for info in infos]
    if COMMENTS not in ordered_names:
        ordered_names.append(COMMENTS)
    result = io.BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in ordered_names:
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, parts[name])
    return result.getvalue()


def _paragraph_index(package: DocxPackage, text: str) -> int:
    paragraphs = list(package.xml(MAIN).iter(w("p")))
    matches = [index for index, paragraph in enumerate(paragraphs) if visible_text(paragraph) == text]
    if len(matches) != 1:
        raise RuntimeError(f"expected one paragraph for {text!r}; got {matches}")
    return matches[0]


def _write_exact(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.read_bytes() != content:
        raise RuntimeError(f"refusing to overwrite different artifact: {path}")
    path.write_bytes(content)


def build() -> dict[str, object]:
    source = _inject_existing_review(_base_docx())
    package = DocxPackage(source)
    replace_index = _paragraph_index(package, "完成访视,记录结果。")
    comment_index = _paragraph_index(package, "给药剂量需医学负责人确认。")

    replace_anchor = make_anchor(source, replace_index, 4, 5)
    comment_anchor = make_anchor(source, comment_index, 0, 4)
    changes = (
        Change(
            "S3-F001",
            replace_anchor,
            "replace",
            after="，",
            rule_id="chinese_sentence_comma",
            user_confirmed=True,
        ),
        Change(
            "S3-F002",
            comment_anchor,
            "comment",
            comment="仅为合成测试：请医学负责人复核给药剂量；本批次不修改该事实。",
            user_confirmed=True,
        ),
    )
    plan = freeze_plan(digest(source), changes, REVIEW_TIME)
    generated = generate(source, plan)

    _write_exact(SOURCE, source)
    _write_exact(OUTPUT, generated.content)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    plan_payload = asdict(plan)
    PLAN.write_text(json.dumps(plan_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = {
        "fixture_scope": "synthetic_only",
        "source_role": "only_modified_object_for_this_synthetic_test",
        "evidence_role": "none",
        "source_path": str(SOURCE.relative_to(ROOT)),
        "source_sha256": generated.source_sha256,
        "output_path": str(OUTPUT.relative_to(ROOT)),
        "output_sha256": generated.output_sha256,
        "plan_path": str(PLAN.relative_to(ROOT)),
        "plan_sha256": generated.plan_sha256,
        "new_revision_ids": list(generated.revision_ids),
        "new_comment_ids": list(generated.comment_ids),
        "review_author": AUTHOR,
        "review_time": REVIEW_TIME,
        "structural_status_at_generation": generated.structural_status,
        "word_visual_status_at_generation": generated.word_visual_status,
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False, indent=2))
